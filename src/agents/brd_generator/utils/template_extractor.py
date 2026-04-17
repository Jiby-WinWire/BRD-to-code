"""Template Extractor for BRD Generator

Extracts template structure from uploaded documents (DOCX, PDF, TXT).
Simplified version without UI dependencies.
"""

import json
import logging
import re
import os
from typing import Dict, Any, Optional, List
from io import BytesIO
from pathlib import Path

logger = logging.getLogger(__name__)


def extract_template_from_document(
    file_path: str,
    template_name: str,
    description: str = ""
) -> Optional[Dict[str, Any]]:
    """Extract template structure from uploaded document
    
    Args:
        file_path: Path to document file (DOCX, PDF, or TXT)
        template_name: Name for the template
        description: Description of the template
        
    Returns:
        Dictionary containing template structure with sections
    """
    file_extension = Path(file_path).suffix.lower()
    
    try:
        # Read file content based on type
        if file_extension == '.pdf':
            content = _extract_from_pdf(file_path)
        elif file_extension in ['.docx', '.doc']:
            content = _extract_from_docx(file_path)
        elif file_extension == '.txt':
            content = _extract_from_txt(file_path)
        else:
            logger.error(f"Unsupported file type: {file_extension}")
            return None
        
        if not content:
            logger.error("No content extracted from document")
            return None
        
        # Parse content to identify sections
        sections = _identify_sections(content)
        
        if not sections:
            logger.warning("No sections identified in document")
            # Create default section
            sections = [{
                "section_id": "content",
                "title": "Content",
                "required": True,
                "description": "Document content",
                "subsections": []
            }]
        
        # Generate template ID from name
        template_id = template_name.lower().replace(' ', '_').replace('-', '_')
        template_id = re.sub(r'[^a-z0-9_]', '', template_id)
        
        # Build template structure
        template = {
            "template_id": template_id,
            "template_name": template_name,
            "description": description or f"Custom template extracted from {Path(file_path).name}",
            "use_case": "Custom user-defined template",
            "sections": sections,
            "formatting": {
                "style": "professional",
                "tone": "business-professional",
                "length": "comprehensive",
                "include_diagrams": False,
                "include_tables": True
            }
        }
        
        logger.info(f"Successfully extracted template with {len(sections)} sections")
        return template
        
    except Exception as e:
        logger.error(f"Error extracting template from document: {str(e)}", exc_info=True)
        return None


def _extract_from_pdf(file_path: str) -> str:
    """Extract text from PDF file"""
    try:
        import PyPDF2
        
        content = []
        with open(file_path, 'rb') as file:
            reader = PyPDF2.PdfReader(file)
            page_count = len(reader.pages)
            
            logger.info(f"Processing PDF with {page_count} pages...")
            
            for i, page in enumerate(reader.pages):
                try:
                    page_text = page.extract_text()
                    if page_text.strip():
                        cleaned_text = page_text.replace('\n\n\n', '\n\n').strip()
                        content.append(f"\n=== PAGE {i+1} ===\n")
                        content.append(cleaned_text + "\n")
                except Exception as e:
                    logger.warning(f"Could not extract text from page {i+1}: {str(e)}")
                    continue
            
            if not content:
                logger.error("No readable text found in PDF")
                return ""
            
            result = "".join(content)
            logger.info(f"Successfully extracted {len(result)} characters from PDF")
            return result
            
    except ImportError:
        logger.error("PyPDF2 library not available. Install with: pip install PyPDF2")
        return ""
    except Exception as e:
        logger.error(f"Error reading PDF file: {str(e)}")
        return ""


def _extract_from_docx(file_path: str) -> str:
    """Extract text from DOCX file"""
    try:
        from docx import Document
        
        doc = Document(file_path)
        content = []
        
        logger.info("Processing Word document...")
        
        # Extract text from paragraphs with style information
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                style_name = paragraph.style.name if paragraph.style else "Normal"
                if 'Heading' in style_name:
                    content.append(f"\n[{style_name.upper()}] {paragraph.text}\n")
                elif paragraph.text.strip().isupper() and len(paragraph.text.strip()) < 100:
                    content.append(f"\n[SECTION] {paragraph.text}\n")
                else:
                    content.append(paragraph.text + "\n")
        
        # Extract text from tables
        for table_idx, table in enumerate(doc.tables, 1):
            content.append(f"\n=== TABLE {table_idx} ===\n")
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells)
                if row_text.strip():
                    content.append(row_text + "\n")
        
        result = "".join(content)
        logger.info(f"Successfully extracted {len(result)} characters from DOCX")
        return result
        
    except ImportError:
        logger.error("python-docx library not available. Install with: pip install python-docx")
        return ""
    except Exception as e:
        logger.error(f"Error reading DOCX file: {str(e)}")
        return ""


def _extract_from_txt(file_path: str) -> str:
    """Extract text from TXT file"""
    try:
        with open(file_path, 'r', encoding='utf-8') as file:
            content = file.read()
        logger.info(f"Successfully read {len(content)} characters from TXT")
        return content
    except Exception as e:
        logger.error(f"Error reading TXT file: {str(e)}")
        return ""


def _identify_sections(content: str) -> List[Dict[str, Any]]:
    """Identify sections in document content
    
    Args:
        content: Document text content
        
    Returns:
        List of section dictionaries
    """
    sections = []
    
    # Common heading patterns
    heading_patterns = [
        r'^[#]+\s+(.+)$',  # Markdown headings
        r'^\[HEADING\d*\]\s*(.+)$',  # Tagged headings from DOCX
        r'^\[SECTION\]\s*(.+)$',  # Tagged sections
        r'^([A-Z][A-Z\s]+)$',  # ALL CAPS headings (short lines)
        r'^(\d+\.?\s+[A-Z].+)$',  # Numbered headings (1. Introduction)
        r'^([A-Z][a-z]+(?:\s+[A-Z][a-z]+)*):?\s*$',  # Title Case headings
    ]
    
    lines = content.split('\n')
    current_section = None
    section_id_counter = 1
    
    for line in lines:
        line = line.strip()
        if not line:
            continue
        
        # Check if line matches any heading pattern
        is_heading = False
        heading_text = None
        
        for pattern in heading_patterns:
            match = re.match(pattern, line)
            if match:
                heading_text = match.group(1).strip()
                # Filter out very long lines (likely not headings)
                if len(heading_text) < 100:
                    is_heading = True
                    break
        
        if is_heading and heading_text:
            # Save previous section
            if current_section:
                sections.append(current_section)
            
            # Create new section
            section_id = heading_text.lower().replace(' ', '_')
            section_id = re.sub(r'[^a-z0-9_]', '', section_id)[:50]  # Limit length
            if not section_id:
                section_id = f"section_{section_id_counter}"
            
            current_section = {
                "section_id": f"{section_id}",
                "title": heading_text,
                "required": True,  # Default to required
                "description": f"{heading_text} section",
                "subsections": []
            }
            section_id_counter += 1
    
    # Add last section
    if current_section:
        sections.append(current_section)
    
    logger.info(f"Identified {len(sections)} sections in document")
    
    # Log section titles for verification
    for section in sections:
        logger.info(f"  - {section['title']}")
    
    return sections


def save_template_to_file(template: Dict[str, Any], output_dir: str = None) -> str:
    """Save extracted template to JSON file
    
    Args:
        template: Template dictionary
        output_dir: Directory to save template (optional)
        
    Returns:
        Path to saved file
    """
    if output_dir is None:
        # Save to templates directory
        current_dir = Path(__file__).parent.parent
        output_dir = current_dir / "templates"
    
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    
    template_id = template.get('template_id', 'custom_template')
    output_file = output_dir / f"{template_id}.json"
    
    with open(output_file, 'w', encoding='utf-8') as f:
        json.dump(template, f, indent=2, ensure_ascii=False)
    
    logger.info(f"Saved template to: {output_file}")
    return str(output_file)


# Example usage
if __name__ == "__main__":
    import sys
    
    if len(sys.argv) < 3:
        print("Usage: python template_extractor.py <file_path> <template_name> [description]")
        sys.exit(1)
    
    file_path = sys.argv[1]
    template_name = sys.argv[2]
    description = sys.argv[3] if len(sys.argv) > 3 else ""
    
    # Extract template
    template = extract_template_from_document(file_path, template_name, description)
    
    if template:
        # Save to file
        output_path = save_template_to_file(template)
        print(f"✅ Template extracted and saved to: {output_path}")
        print(f"📋 Sections found: {len(template['sections'])}")
        for section in template['sections']:
            print(f"   - {section['title']}")
    else:
        print("❌ Failed to extract template from document")
        sys.exit(1)
